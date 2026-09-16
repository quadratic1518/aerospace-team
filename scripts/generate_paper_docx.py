"""Generate RESEARCH_PAPER.docx from RESEARCH_PAPER.md."""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "RESEARCH_PAPER.md"
DOCX_PATH = ROOT / "RESEARCH_PAPER.docx"


def add_formatted_paragraph(doc, text, style=None, bold=False, italic=False, size=11):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    return p


def parse_inline(text, paragraph):
    """Handle **bold** and *italic* inline formatting."""
    parts = re.split(r"(\*\*.*?\*\*|\*.*?\*|`.*?`|\$.*?\$)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)
        for run in paragraph.runs:
            if run.font.name != "Courier New":
                run.font.name = "Times New Roman"
            if not run.font.size:
                run.font.size = Pt(11)


def convert_md_to_docx(md_path: Path, docx_path: Path):
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    i = 0
    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if in_code_block:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Courier New"
                run.font.size = Pt(9)
                p.paragraph_format.left_indent = Inches(0.3)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        if line.strip() == "---":
            doc.add_paragraph()
            i += 1
            continue

        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue

        if line.startswith("## "):
            p = doc.add_heading(line[3:].strip(), level=1)
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue

        if line.startswith("### "):
            p = doc.add_heading(line[4:].strip(), level=2)
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue

        if line.startswith("#### "):
            p = doc.add_heading(line[5:].strip(), level=3)
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue

        if "|" in line and line.strip().startswith("|"):
            if not in_table:
                in_table = True
                table_rows = []
            if re.match(r"^\|[\s\-:|]+\|$", line.strip()):
                i += 1
                continue
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            table_rows.append(cells)
            if i + 1 >= len(lines) or "|" not in lines[i + 1]:
                if table_rows:
                    ncols = max(len(r) for r in table_rows)
                    table = doc.add_table(rows=len(table_rows), cols=ncols)
                    table.style = "Table Grid"
                    for ri, row in enumerate(table_rows):
                        for ci in range(ncols):
                            val = row[ci] if ci < len(row) else ""
                            val = re.sub(r"\*\*(.*?)\*\*", r"\1", val)
                            table.rows[ri].cells[ci].text = val
                    doc.add_paragraph()
                in_table = False
                table_rows = []
            i += 1
            continue

        if line.strip().startswith("- ") or line.strip().startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            parse_inline(line.strip()[2:], p)
            i += 1
            continue

        if re.match(r"^\d+\.\s", line.strip()):
            p = doc.add_paragraph(style="List Number")
            parse_inline(re.sub(r"^\d+\.\s", "", line.strip()), p)
            i += 1
            continue

        if line.strip().startswith("$$") or (line.strip().startswith("$") and line.strip().endswith("$")):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line.strip().strip("$"))
            run.italic = True
            run.font.name = "Times New Roman"
            i += 1
            continue

        if line.strip() == "":
            i += 1
            continue

        if line.strip().startswith("```") or line.strip().startswith("┌") or line.strip().startswith("│") or line.strip().startswith("└") or line.strip().startswith("─") or line.strip().startswith("•"):
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            i += 1
            continue

        p = doc.add_paragraph()
        parse_inline(line.strip(), p)
        i += 1

    doc.save(docx_path)
    print(f"Generated: {docx_path}")


if __name__ == "__main__":
    convert_md_to_docx(MD_PATH, DOCX_PATH)
